"""
Build agent-prompt mapping spreadsheet across three sources:
  1. Prompts/system-prompts.json  (UI catalog)
  2. agents/*.yaml / AGENT_DEFINITIONS.md  (contract specs)
  3. backend/src/services/*.ts + backend/bakeoff-roadmap.ts  (hardcoded runtime)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# DATA  (collected by manual read of all three sources)
# ---------------------------------------------------------------------------

# Each entry: (group_label, sp_name, sp_id, sp_content, ag_file, ag_id, ag_prompt, svc_file, svc_const, svc_content, match_quality)
ROWS = [
    (
        "Engineer /\nCodeGenerator",
        "Engineer System Prompt",
        "prompt-engineer-system",
        "You are a senior software engineer responsible for full functional implementation, testing, and validation. Based on the Researcher's recommendations, produce:\n1. Detailed architecture (frontend, backend, build pipeline).\n2. Feature-by-feature plan ensuring 100% coverage of the goal requirements.\n3. Validation and testing strategy for every subsystem.\n4. Error handling, data integrity, and CI/CD considerations.\n5. Example test cases or scripts that confirm correctness.\nBe exhaustive and pragmatic.",
        "agents/engineer.yaml\n(ag_20251202_engineer)",
        "ag_20251202_engineer",
        "Engineer. Produce complete, working implementation files.\n\nFor every file: ## File: path/to/file.ext\\n```lang\\ncontent```.\nNo prose before or after file blocks unless the task explicitly asks for explanation.\n\nRules:\n- Minimal assumptions. State any that are load-bearing in a comment inside the file.\n- Error handling at system boundaries (I/O, network, user input).\n- Use the stack language and versions specified in the task.\n- Smallest correct implementation first — no speculative abstractions.",
        "backend/src/services/task-executor.ts\nbackend/bakeoff-roadmap.ts",
        "FALLBACK_SYSTEM_PROMPT\nROLE_PROMPTS.engineer",
        "FALLBACK: 'You are a software engineer. Execute the assigned task and produce detailed, production-ready output. Follow the provided project stack constraints when they exist. For code tasks, write complete, runnable code.'\n\nBAKEOFF: 'Engineer. Produce complete, working implementation files.\\nFor every file: ## File: path/to/file.ext\\n```lang\\ncontent```. No prose before/after blocks.\\nMinimal correct implementation, error handling at boundaries, strict TS, no runtime deps.'",
        "SIMILAR\nagents/ yaml prompt ≈ bakeoff inline;\nUI catalog is older/more verbose",
    ),
    (
        "Researcher",
        "Researcher System Prompt",
        "prompt-researcher-system",
        "You are a senior technical researcher. Analyze the user's goal and propose at least three distinct architectural strategies for achieving it. For each, outline tradeoffs, scalability, and privacy implications. Include possible frameworks, libraries, and deployment models.",
        "agents/researcher.yaml\n(ag_20251109_researcher)",
        "ag_20251109_researcher",
        "Researcher. Identify facts, options, and risks that unblock implementation.\n\nOutput:\n  FACTS: bullet list — constraints, options (pros/cons), key decisions needed\n  RISKS: Risk | Impact | Mitigation (one per line)\n\nNo hype, no fabrication. Label assumptions. Think: what facts unblock the next build decision?",
        "(none)",
        "(none)",
        "Not used in any backend service or bakeoff arm. Research role is absorbed into the planner/architect pipeline.",
        "NO MATCH\nin backend services",
    ),
    (
        "Critic",
        "Critic System Prompt",
        "prompt-critic-system",
        "You are a technical reviewer and performance analyst. Evaluate the Engineer's plan for maintainability, performance, scalability, and risk. Highlight potential issues, oversights, and suggest refinements or alternative methods to reduce complexity.",
        "agents/critic.yaml\n(ag_20251202_critic)",
        "ag_20251202_critic",
        "Critic. Find defects, security risks, and correctness issues.\n\nOutput findings as:\n  severity(low|medium|high|critical) | file:line | issue | fix\n\nFor corrected code: ## File: path/to/file.ext\\n```lang\\ncontent```.\nNo speculative issues — verifiable problems only, with a specific fix for each.",
        "backend/bakeoff-roadmap.ts",
        "ROLE_PROMPTS.critic",
        "'Critic. Review existing files for defects and contract gaps. Where a fix is needed, output the corrected file as a ## File: block. Strict TS. No speculative changes.'",
        "SIMILAR\nagents/ yaml ≈ bakeoff inline;\nUI catalog is diffuse/non-actionable",
    ),
    (
        "Synthesizer",
        "Synthesizer System Prompt",
        "prompt-synthesizer-system",
        "You are an integration specialist and communicator. Combine all prior outputs into a cohesive, actionable implementation plan. Maintain clarity and conciseness. Include an ordered execution roadmap (Phase 1 → Phase 3) and note dependencies or resource needs.",
        "agents/synthesizer.yaml\n(ag_20251202_synthesizer)",
        "ag_20251202_synthesizer",
        "Synthesizer. Merge analysis, implementation, and critique into a coherent execution plan.\n\nResolve contradictions explicitly: Conflict | Decision | Rationale.\nPrefer shippable slices — MVP first, hardening second.\nOutput phases with: objective, deliverables, acceptance criteria.\nFor any files produced: ## File: path/to/file.ext\\n```lang\\ncontent```.",
        "(none)",
        "(none)",
        "Not used in any backend service or bakeoff arm. Synthesis step is implicit in the planner output.",
        "NO MATCH\nin backend services",
    ),
    (
        "Commissioner",
        "Commissioner System Prompt",
        "prompt-commissioner-system",
        "You are The Commissioner — the pragmatic evaluator. Your job is to assess real-world value, usability, and empowerment potential. Ask: Does this design meaningfully help solve the problem? Will it save time and resources? Assign a Value Score (0-10) and explain your reasoning. If the score is below 7, recommend concrete adjustments.",
        "agents/commissioner.yaml\n(ag_20251202_commissioner)",
        "ag_20251202_commissioner",
        "You are the Commissioner.\nGoal: decide whether the synthesized plan is worth doing *now*, and under what conditions.\nRules:\n- Be blunt. If it's not worth doing, say so.\n- Optimize for ROI, risk, opportunity cost, user impact.\n- Prefer 'conditional go' with crisp conditions over vague hedging.\nScoring: 9-10 = high ROI; 7-8 = solid; 4-6 = unclear; 0-3 = not viable.\nOutput valid JSON: {value_score, recommendation, rationale, conditions, improvements}",
        "(none)",
        "(none)",
        "Not used in any backend service or bakeoff arm. Approval gates are handled by the human-driven approval-service.ts (not an LLM call).",
        "NO MATCH\nin backend services",
    ),
    (
        "Supervisor",
        "Supervisor System Prompt",
        "prompt-supervisor-system",
        "You are the Supervisor. Score the quality of the current run output (0-10), issue corrective guidance for any deficiencies, and promote durable insights to global memory. Ensure all agents met their objectives and the final deliverable is production-ready.",
        "agents/supervisor.yaml\n(ag_20251202_supervisor)",
        "ag_20251202_supervisor",
        "You are the Supervisor.\nGoal: judge orchestration quality, give corrective guidance, and extract reusable learnings.\nRules:\n- Score outputs on: completeness, correctness, usability, efficiency, fit-to-goal.\n- Provide feedback as concrete actions (not motivational quotes).\n- Extract durable 'playbook' insights: what to repeat, what to avoid.\nOutput valid JSON: {quality_score, feedback[], insights[], agent_scores{}}",
        "(none)",
        "(none)",
        "Not used in any backend service or bakeoff arm. Quality supervision is a future capability (Phase 40+).",
        "NO MATCH\nin backend services",
    ),
    (
        "Architect /\nDesignPlanGenerator",
        "Design Plan Generation Prompt",
        "prompt-design-plan-generation",
        "Given the approved application intake, generate a detailed design plan that includes:\n1. High-level architecture (components, data flow, interfaces)\n2. Technology stack recommendations with justification\n3. Database schema outline\n4. API endpoint definitions\n5. Frontend component hierarchy\n6. Phase-by-phase implementation roadmap\nEnsure each phase has clear deliverables and acceptance criteria.",
        "agents/architect_agent.yaml\n(ag_20260407_architect)\n+AGENT_DEFINITIONS.md: ArchitectureDesigner",
        "ag_20260407_architect",
        "YAML: You are Architect Agent. Convert approved intent into an implementation-ready architecture that is minimal, resilient, and verifiable. Present ≥2 viable options; select one with explicit trade-offs. Define boundaries, contracts, invariants, and failure handling. Required output: Context+constraints, Option analysis, Target design, Delivery plan, Validation hooks.\n\nAGENT_DEFINITIONS ArchitectureDesigner: You are a principal systems architect. Name the tech stack EXPLICITLY. List 3-5 core components each with Name, Responsibility, Input/Output interfaces. Be opinionated. Output must be valid JSON.",
        "backend/src/services/design-plan-generator.ts\nbackend/bakeoff-roadmap.ts",
        "SYSTEM_PROMPT\nROLE_PROMPTS.architect",
        "DESIGN_PLAN: 'You are an expert software architect. Given an application description, produce a thorough design plan as a single JSON object. Schema: {overview, architecture, components[{name,description,responsibility,interfaces[]}], tradeoffs[], recommendations[]}. Rules: components must cover ALL aspects; tradeoffs must be honest; respond with ONLY the JSON object.'\n\nBAKEOFF: 'Architect. Produce design + type/schema files only: TypeScript interfaces, typed registries, JSON-shaped schema objects, route maps. Output ## File: blocks. No prose. Strict TS.'",
        "SIMILAR\nAll three target architecture output;\nUI catalog is generic prose;\nagents/ yaml is richer process;\nbackend is JSON-schema-enforced",
    ),
    (
        "RoadmapPlanner",
        "(none — closest is\n'Design Plan Generation\nPrompt' but that's Architect)",
        "(none)",
        "(none)",
        "AGENT_DEFINITIONS.md:\nRoadmapPlanner",
        "(in AGENT_DEFINITIONS.md only)",
        "AGENT_DEFINITIONS RoadmapPlanner: You are an experienced project manager. Each task must specify EXACTLY what files it produces. Format: 'Create file src/services/items.ts with: [brief description]'. Tasks are atomic: one task = one deliverable. Dependencies are explicit. Do NOT create 'setup' or 'planning' tasks. Phases: Foundation → Core Features → Testing → Hardening → Deployment. Output valid JSON.",
        "backend/src/services/roadmap-generator.ts\nbackend/bakeoff-roadmap.ts",
        "SYSTEM_PROMPT\nPLANNER_PROMPT",
        "ROADMAP_GENERATOR: 'Generate an implementation roadmap as a single JSON object. Schema: {title, description, estimatedDuration, phases[{id,number,name,goal,estimatedHours,dependencies,tasks[{id,name,description(compact directive: ACTION|OUTPUT|CONSTRAINT),status,estimatedHours,dependencies}]}]}. 3-6 phases, 3-6 tasks per phase. Respond with ONLY the JSON.'\n\nBAKEOFF PLANNER: 'You are the Roadmap Builder + Supervisor. Decompose ONLY the First Build Slice (core skeleton) into an execution plan as a single JSON. Schema: {phases[{name,tasks[{name,description,agent:architect|engineer|test|critic}]}]}. 3-5 phases, 2-4 tasks per phase. Name every file each task outputs.'",
        "SIMILAR\nNo UI catalog entry;\nagents/ has no YAML file (def only);\nboth backend variants produce JSON phases+tasks",
    ),
    (
        "Test &\nValidation",
        "(none)",
        "(none)",
        "(none)",
        "agents/test_and_validation_agent.yaml\n(ag_20260407_test_validation)",
        "ag_20260407_test_validation",
        "Test and Validation Agent. Write complete test files and report results.\n\nFor every test file: ## File: path/to/file.test.ts\\n```typescript\\ncontent```.\nUse describe/it/expect (vitest). Cover: happy path, edge cases, invalid input.\nNo prose — only file blocks and, if needed, a brief pass/fail summary at the end.",
        "backend/bakeoff-roadmap.ts",
        "ROLE_PROMPTS.test",
        "'Test and Validation Agent. Write complete vitest test files.\\nFor every file: ## File: path/to/file.test.ts\\n```typescript\\ncontent```. Use describe/it/expect imported from 'vitest'. Cover happy path, edges, invalid input. No prose — only file blocks.'",
        "SIMILAR\nNo UI catalog entry;\nagents/ yaml ≈ bakeoff inline (nearly identical)",
    ),
    (
        "RepairOrchestrator",
        "(none)",
        "(none)",
        "(none)",
        "AGENT_DEFINITIONS.md:\nRepairOrchestrator",
        "(in AGENT_DEFINITIONS.md only)",
        "AGENT_DEFINITIONS RepairOrchestrator: You are a debugging expert. A code generation task failed validation. Analyze errors from tsc/vitest. Categorize: missing import, type mismatch, syntax, logic. Determine which agent to invoke (CodeGenerator or ArchitectureDesigner). CRITICAL: If same error signature in previousRepairAttempts, ESCALATE. Max 3 attempts per task. Do NOT attempt repairs yourself. Delegate.",
        "backend/src/services/repair-strategist.ts\nbackend/src/services/repair-task-builder.ts\nbackend/src/services/repair-policy.ts",
        "(pure logic,\nno LLM system prompt)",
        "Repair is handled as pure TypeScript logic:\n- repair-strategist.ts: classifies failure type and picks strategy (repair / escalate / skip)\n- repair-task-builder.ts: groups tsc/vitest errors by file into targeted RepairTask objects\n- repair-policy.ts: enforces max-retry circuit breaker\nNo LLM system prompt is used for repair routing — the agent-selector routes repair tasks to the best-fit agent based on task keywords.",
        "DIFFERENT\nUI: no entry;\nagents/: spec only (no YAML);\nbackend: pure logic, no LLM prompt",
    ),
    (
        "Application\nIntake",
        "Application Intake Prompt",
        "prompt-application-intake",
        "You are an expert software architect reviewing an application intake request. Extract the following from the user's description:\n1. Application name and purpose\n2. Core features (numbered list)\n3. Target users\n4. Technical constraints or preferences\n5. Success criteria\nReturn structured JSON matching the ApplicationIntake schema.",
        "(none — no YAML\nin agents/ dir)",
        "(none)",
        "(none)",
        "(none — intake handled\nby HTTP route directly)",
        "(none)",
        "Application creation is a direct POST /api/applications route that persists the user-submitted form data. No LLM call is made at intake; LLM calls begin at the design-plan generation step.",
        "NO MATCH\nin agents/ or backend services;\nUI catalog only",
    ),
]

# ---------------------------------------------------------------------------
# BUILD WORKBOOK
# ---------------------------------------------------------------------------

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Agent Prompt Map"

# -- Styles --
HDR_FILL   = PatternFill("solid", fgColor="1F3864")  # dark navy
GRP_FILL   = PatternFill("solid", fgColor="D6E4F0")  # light blue
SAME_FILL  = PatternFill("solid", fgColor="C6EFCE")  # green
SIM_FILL   = PatternFill("solid", fgColor="FFEB9C")  # yellow
DIFF_FILL  = PatternFill("solid", fgColor="FFC7CE")  # red
NONE_FILL  = PatternFill("solid", fgColor="F2F2F2")  # grey

HDR_FONT   = Font(bold=True, color="FFFFFF", size=10)
GRP_FONT   = Font(bold=True, size=10)
BODY_FONT  = Font(size=9)
WRAP       = Alignment(wrap_text=True, vertical="top")
CENTER     = Alignment(horizontal="center", wrap_text=True, vertical="top")

thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def cell(ws, r, c, value, fill=None, font=None, align=None):
    ws.cell(row=r, column=c, value=value)
    if fill:  ws.cell(row=r, column=c).fill  = fill
    if font:  ws.cell(row=r, column=c).font  = font
    if align: ws.cell(row=r, column=c).alignment = align
    ws.cell(row=r, column=c).border = BORDER

HEADERS = [
    "Agent / Role Group",
    "system-prompts.json\nName",
    "system-prompts.json\nID",
    "system-prompts.json\nPrompt Content",
    "agents/ Directory\nFile + Agent ID",
    "agents/ Directory\nAgent ID",
    "agents/ Directory\nPrompt Content",
    "backend/src/services\nFile(s)",
    "backend/src/services\nConst Name",
    "backend/src/services\nHardcoded Prompt",
    "Match Quality",
]

COL_WIDTHS = [18, 20, 28, 55, 28, 28, 55, 32, 22, 55, 20]

# Header row
for c, (hdr, w) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
    cell(ws, 1, c, hdr, fill=HDR_FILL, font=HDR_FONT, align=CENTER)
    ws.column_dimensions[get_column_letter(c)].width = w

ws.row_dimensions[1].height = 40

# Data rows
for i, row in enumerate(ROWS, start=2):
    (group, sp_name, sp_id, sp_content,
     ag_file, ag_id, ag_prompt,
     svc_file, svc_const, svc_content,
     match) = row

    mq = match.upper()
    if "SAME" in mq:
        mfill = SAME_FILL
    elif "SIMILAR" in mq:
        mfill = SIM_FILL
    elif "DIFFERENT" in mq:
        mfill = DIFF_FILL
    else:
        mfill = NONE_FILL

    no_fill = NONE_FILL if sp_name == "(none)" or sp_name.startswith("(none") else None

    cell(ws, i,  1, group,       fill=GRP_FILL, font=GRP_FONT, align=WRAP)
    cell(ws, i,  2, sp_name,     fill=no_fill,  font=BODY_FONT, align=WRAP)
    cell(ws, i,  3, sp_id,       fill=no_fill,  font=BODY_FONT, align=WRAP)
    cell(ws, i,  4, sp_content,  fill=no_fill,  font=BODY_FONT, align=WRAP)
    cell(ws, i,  5, ag_file,     font=BODY_FONT, align=WRAP)
    cell(ws, i,  6, ag_id,       font=BODY_FONT, align=WRAP)
    cell(ws, i,  7, ag_prompt,   font=BODY_FONT, align=WRAP)
    cell(ws, i,  8, svc_file,    font=BODY_FONT, align=WRAP)
    cell(ws, i,  9, svc_const,   font=BODY_FONT, align=WRAP)
    cell(ws, i, 10, svc_content, font=BODY_FONT, align=WRAP)
    cell(ws, i, 11, match,       fill=mfill,    font=Font(bold=True, size=9), align=WRAP)

    ws.row_dimensions[i].height = 120

# Freeze header
ws.freeze_panes = "B2"

out = r"f:\Development\20_Staging\AI Projects\2026-06-13_Orchestration\docs\agent_prompt_map.xlsx"
wb.save(out)
print(f"Saved: {out}")
